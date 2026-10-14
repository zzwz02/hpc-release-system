"""Weekly report check — workbook parsing, date rules and problem decisions.

Ported from the local weekly_report_jira_check.py; the rules are unchanged.
Pure functions only: no DB, JIRA or LLM access.  The service
(app/services/weekly_check_service.py) feeds JIRA activity and LLM replies in.

Workbook layout: one sheet per person, column A the date, column B the text.
The first dated row from the top is the newest week; its column B is
"本周更新", cut where "下周计划" starts.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass
from typing import Any

T_BAD_DATE = "周报日期异常"
T_NO_REPORT = "周报未更新"
T_NO_JIRA = "未写JIRA"
T_BAD_ID = "JIRA ID不完整"
T_NOT_FOUND = "JIRA不存在"
T_FORBIDDEN = "JIRA无权限"
T_NOT_UPDATED = "JIRA未更新"
T_MISMATCH = "更新与周报不符"
T_PARTIAL = "更新与周报部分相符"
# Column order of the summary table and of the problem-type filter.
PROBLEM_TYPES = (
    T_BAD_DATE, T_NO_REPORT, T_NO_JIRA, T_BAD_ID, T_NOT_FOUND,
    T_FORBIDDEN, T_NOT_UPDATED, T_MISMATCH, T_PARTIAL,
)

# A top date before this year is a typo, not an old report.
MIN_REPORT_YEAR = 2025
# A top date more than this many days past today is a typo.
MAX_FUTURE_DAYS = 7
# A top date more than this many days before the report date means the sheet
# still shows an earlier week (or its date was not changed).
STALE_REPORT_DAYS = 3
# JIRA text handed to the consistency prompt per issue.
UPDATE_TEXT_LIMIT = 3000

JIRA_KEY_RE = re.compile(r"^[A-Z][A-Z0-9]+-\d+$")
_PLAN_RE = re.compile(r"下周计划|下周[:：]")
_DATE_FORMATS = ("%Y-%m-%d", "%m-%d-%y", "%Y/%m/%d", "%m/%d/%Y")


@dataclass(frozen=True)
class SheetReport:
    """The newest entry of one sheet."""

    sheet: str
    date: dt.date
    newest: dt.date
    text: str


def problem(kind: str, item: str, jira_key: str = "", **extra: str) -> dict[str, str]:
    """One finding; keys match the weekly_check_findings columns."""
    return {
        "problem": kind,
        "item": item,
        "jira_key": jira_key,
        "jira_status": extra.get("jira_status", ""),
        "assignee": extra.get("assignee", ""),
        "last_updated": extra.get("last_updated", ""),
        "last_updater": extra.get("last_updater", ""),
        "note": extra.get("note", ""),
    }


# ---------------------------------------------------------------------------
# Workbook parsing
# ---------------------------------------------------------------------------

def to_date(value: object) -> dt.date | None:
    """Return the date a column-A cell holds, or None when it is not a date."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for fmt in _DATE_FORMATS:
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def cut_plan(text: str) -> str:
    """Drop "下周计划" and everything after it."""
    match = _PLAN_RE.search(text)
    return text[: match.start()] if match else text


def read_sheet(name: str, rows: list[tuple[object, object]]) -> SheetReport | None:
    """Build a SheetReport from (column A, column B) data rows; None without a dated row."""
    top: tuple[dt.date, str] | None = None
    dates: list[dt.date] = []
    for first, second in rows:
        found = to_date(first)
        if not found:
            continue
        dates.append(found)
        if top is None:
            top = (found, cut_plan(str(second or "")).strip())
    if top is None:
        return None
    return SheetReport(sheet=name, date=top[0], newest=max(dates), text=top[1])


def read_reports(content: bytes) -> tuple[list[SheetReport], list[str]]:
    """Parse an .xlsx held in memory.

    Returns one SheetReport per sheet that has a dated row, plus the names of
    the sheets that have none.  Raises ValueError for an unreadable workbook.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("服务器未安装 openpyxl：请运行 python -m pip install openpyxl") from exc
    try:
        # read_only streams cell values; macros and external links never run.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取该文件：不是有效的 .xlsx 周报") from exc
    try:
        reports: list[SheetReport] = []
        undated: list[str] = []
        for sheet in workbook.worksheets:
            name = sheet.title.strip()
            rows = [
                tuple((list(row) + [None, None])[:2])
                for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True)
            ]
            report = read_sheet(name, rows)
            if report is None:
                undated.append(name)
            else:
                reports.append(report)
        return reports, undated
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# Report date and week
# ---------------------------------------------------------------------------

def parse_report_date(value: str) -> dt.date:
    try:
        return dt.datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"报告日期格式应为 YYYY-MM-DD：{value!r}") from exc


def default_report_date(reports: list[SheetReport]) -> dt.date | None:
    """The most common top date, so one person a day early or late does not matter; a tie goes to the later date."""
    counts = Counter(report.date for report in reports if report.date.year >= MIN_REPORT_YEAR)
    if not counts:
        return None
    return max(counts, key=lambda day: (counts[day], day))


def week_of(report_date: dt.date) -> dict[str, str]:
    """ISO week of a report date: key ``2026-W41`` plus its Monday and Sunday."""
    year, week, weekday = report_date.isocalendar()
    start = report_date - dt.timedelta(days=weekday - 1)
    end = start + dt.timedelta(days=6)
    return {
        "week_key": f"{year}-W{week:02d}",
        "week_start": start.isoformat(),
        "week_end": end.isoformat(),
        "label": f"{year} W{week:02d}（{start.month}月{start.day}日-{end.month}月{end.day}日）",
    }


# ---------------------------------------------------------------------------
# Date rules
# ---------------------------------------------------------------------------

def date_problems(report: SheetReport, *, report_date: dt.date, today: dt.date) -> list[dict[str, str]]:
    """Date problems of one sheet.  They are reported; the entry is still analysed."""
    top = report.date
    problems = []
    if top.year < MIN_REPORT_YEAR or top > today + dt.timedelta(days=MAX_FUTURE_DAYS):
        problems.append(problem(T_BAD_DATE, f"最上面一条日期为 {top}，日期明显错误"))
    elif report.newest > top:
        problems.append(problem(T_BAD_DATE, f"最上面一条日期为 {top}，不是最新（表中最新为 {report.newest}）"))
    elif top < report_date - dt.timedelta(days=STALE_REPORT_DAYS):
        # Either this week was not written, or its date was not changed.
        problems.append(problem(T_NO_REPORT, f"最上面一条日期为 {top}"))
    if not report.text:
        problems.append(problem(T_NO_REPORT, "本周更新为空"))
    return problems


# ---------------------------------------------------------------------------
# Work items and JIRA activity
# ---------------------------------------------------------------------------

def text_sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def item_keys(item: dict[str, Any]) -> list[str]:
    return [str(key).strip().upper() for key in item.get("jira_ids") or []]


def updated_by(activity: dict[str, Any], username: str, since: dt.datetime) -> bool:
    """True when the person changed a field, commented, logged work or created the issue since *since*."""
    return any(event["time"] >= since and event["user"] == username for event in activity["events"])


def jira_update_text(
    activity: dict[str, Any],
    username: str,
    since: dt.datetime,
    limit: int = UPDATE_TEXT_LIMIT,
) -> str:
    """The person's updates to one issue inside the window, as text for the consistency prompt."""
    lines = []
    for change in activity["changes"]:
        if change["user"] == username and change["time"] >= since:
            stamp = change["time"].strftime("%m-%d")
            for entry in change["items"]:
                lines.append(f"{stamp} 修改 {entry['field']}：{entry['from'][:200]} → {entry['to'][:500]}")
    for comment in activity["comments"]:
        if username in (comment["author"], comment["update_author"]) and comment["time"] >= since:
            lines.append(f"{comment['time'].strftime('%m-%d')} comment：{comment['body'].strip()[:1500]}")
    for worklog in activity["worklogs"]:
        if worklog["author"] == username and worklog["time"] >= since:
            lines.append(
                f"{worklog['time'].strftime('%m-%d')} worklog {worklog['time_spent']}：{worklog['comment'][:500]}"
            )
    return (f"[{activity['key']}] {activity['summary']}\n" + "\n".join(lines))[:limit]


def item_problems(
    item: dict[str, Any],
    *,
    username: str,
    issues: dict[str, dict[str, Any] | str],
    since: dt.datetime,
) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    """Decide one work item's problems.

    *issues* maps each JIRA key to its activity, or to T_NOT_FOUND / T_FORBIDDEN
    when JIRA refused it.  Without *username* only the id problems are decided.
    Returns the problems and the person's own issues that were updated in the
    window (the ones a consistency check compares with the report).

    Only issues assigned to the person are checked: one assigned to somebody
    else may have been filed for them.  Closed / Resolved issues are checked
    too.  An item with several own issues counts as updated when any one is.
    """
    title = str(item.get("title") or "")
    keys = item_keys(item)
    bare = item.get("incomplete_ids") or []
    if not keys and not bare:
        return [problem(T_NO_JIRA, title)], []
    problems = [problem(T_BAD_ID, title, str(number)) for number in bare]
    if not username:
        return problems, []
    stale: list[dict[str, str]] = []
    updated: list[dict[str, Any]] = []
    for key in keys:
        if not JIRA_KEY_RE.match(key):
            problems.append(problem(T_BAD_ID, title, key))
            continue
        activity = issues[key]
        if isinstance(activity, str):
            problems.append(problem(activity, title, key))
            continue
        assignee = activity["assignee"] or {}
        if assignee.get("name") != username:
            continue
        if updated_by(activity, username, since):
            updated.append(activity)
            continue
        last = max(activity["events"], key=lambda event: event["time"]) if activity["events"] else None
        stale.append(problem(
            T_NOT_UPDATED, title, key,
            jira_status=activity["status"],
            assignee=assignee.get("display_name", ""),
            last_updated=last["time"].strftime("%Y-%m-%d %H:%M") if last else "",
            last_updater=(last["display"] or last["user"]) if last else "",
        ))
    if not updated:
        problems.extend(stale)
    return problems, updated


# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------

def summarize(findings: list[dict[str, Any]]) -> dict[str, Any]:
    by_problem = Counter(row["problem"] for row in findings)
    by_person = Counter(row["person"] for row in findings)
    return {
        "total": len(findings),
        "by_problem": [
            {"problem": kind, "count": by_problem[kind]}
            for kind in PROBLEM_TYPES if by_problem.get(kind)
        ],
        "by_person": [
            {"person": person, "count": count}
            for person, count in sorted(by_person.items(), key=lambda pair: (-pair[1], pair[0].lower()))
        ],
    }
